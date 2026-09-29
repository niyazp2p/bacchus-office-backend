import io
import uuid
from datetime import datetime, timezone
from typing import Annotated
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, status, Response
from fastapi.responses import StreamingResponse
from sqlalchemy import select, func, desc, or_
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl
from app.services.pdf_engine import generate_lead_dossier_pdf

from app.core.database import get_db
from app.models.office_auth import OfficeUser, OfficeEmployee, OfficeRole
from app.models.office_lead import OfficeLead, LeadTier, LeadStatus, CommercialModel
from app.schemas.office_lead import LeadCreate, LeadUpdate, LeadPatchStatus, LeadResponse, RoutingResult
from app.dependencies.office_auth import get_current_user, require_roles
from app.services.lead_engine import generate_dedup_hash, calculate_lead_score

router = APIRouter(prefix="/leads", tags=["Office Leads Ingestion & Routing Engine"])

async def generate_lead_code(db: AsyncSession) -> str:
    stmt = select(func.count(OfficeLead.id))
    result = await db.execute(stmt)
    count = (result.scalar() or 0) + 1
    suffix = uuid.uuid4().hex[:4].upper()
    return f"BAC-LD-{count:05d}-{suffix}"

# 1. Manual Lead Ingestion (Single)
@router.post("", response_model=LeadResponse, status_code=status.HTTP_201_CREATED)
async def create_manual_lead(
    payload: LeadCreate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[OfficeUser, Depends(get_current_user)],
):
    dedup_hash = generate_dedup_hash(payload.email, payload.phone)
    existing = await db.scalar(select(OfficeLead).where(OfficeLead.dedup_hash == dedup_hash))
    if existing:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"Duplicate lead detected! Lead already exists with code: {existing.lead_code}",
        )

    score, tier = calculate_lead_score(
        commercial_model=payload.commercial_model,
        volume_estimate=payload.volume_estimate,
        country=payload.country,
        email=payload.email,
        phone=payload.phone,
    )
    lead_code = await generate_lead_code(db)

    lead = OfficeLead(
        lead_code=lead_code,
        company_name=payload.company_name.strip(),
        contact_name=payload.contact_name.strip(),
        email=payload.email.strip().lower(),
        phone=payload.phone.strip() if payload.phone else None,
        country=payload.country.strip(),
        state=payload.state.strip() if payload.state else None,
        commercial_model=payload.commercial_model,
        volume_estimate=payload.volume_estimate.strip() if payload.volume_estimate else None,
        score=score,
        tier=tier,
        status=LeadStatus.NEW,
        dedup_hash=dedup_hash,
    )
    db.add(lead)
    await db.commit()
    await db.refresh(lead)
    return lead

# 2. Bulk Excel Upload Ingestion
@router.post("/bulk-import")
async def bulk_import_leads_excel(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: OfficeUser = Depends(require_roles([OfficeRole.SUPER_ADMIN, OfficeRole.COO, OfficeRole.HR_MANAGER])),
):
    if not file.filename.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Invalid file type. Only .xlsx files are supported.")

    contents = await file.read()
    workbook = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    sheet = workbook.active

    imported_count = 0
    duplicate_count = 0
    errors = []

    # Expected Header: [company_name, contact_name, email, phone, country, state, commercial_model, volume_estimate]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    for idx, row in enumerate(rows, start=2):
        if not row or not any(row):
            continue
        try:
            company_name = str(row[0] or "").strip()
            contact_name = str(row[1] or "").strip()
            email = str(row[2] or "").strip().lower()
            phone = str(row[3] or "").strip() if row[3] else None
            country = str(row[4] or "India").strip()
            state = str(row[5] or "").strip() if row[5] else None
            model_raw = str(row[6] or "DISTRIBUTION").strip().upper()
            volume = str(row[7] or "").strip() if len(row) > 7 and row[7] else None

            if not company_name or not email:
                continue

            commercial_model = CommercialModel.DISTRIBUTION
            if model_raw in CommercialModel.__members__:
                commercial_model = CommercialModel[model_raw]

            dedup = generate_dedup_hash(email, phone)
            exists = await db.scalar(select(OfficeLead.id).where(OfficeLead.dedup_hash == dedup))
            if exists:
                duplicate_count += 1
                continue

            score, tier = calculate_lead_score(commercial_model, volume, country, email, phone)
            lead_code = await generate_lead_code(db)

            new_lead = OfficeLead(
                lead_code=lead_code,
                company_name=company_name,
                contact_name=contact_name,
                email=email,
                phone=phone,
                country=country,
                state=state,
                commercial_model=commercial_model,
                volume_estimate=volume,
                score=score,
                tier=tier,
                status=LeadStatus.NEW,
                dedup_hash=dedup,
            )
            db.add(new_lead)
            imported_count += 1
        except Exception as e:
            errors.append(f"Row {idx}: {str(e)}")

    await db.commit()
    return {
        "success": True,
        "imported_records": imported_count,
        "duplicates_skipped": duplicate_count,
        "errors": errors[:10],
    }

# 3. Automated Load-Balanced Routing to Telecallers
@router.post("/route-to-callers", response_model=RoutingResult)
async def route_leads_to_active_callers(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[OfficeUser, Depends(require_roles([OfficeRole.SUPER_ADMIN, OfficeRole.COO, OfficeRole.OPERATIONS_HEAD]))],
):
    # Fetch all active Telecaller employees
    callers_stmt = (
        select(OfficeEmployee)
        .join(OfficeUser, OfficeEmployee.user_id == OfficeUser.id)
        .where(OfficeUser.role == OfficeRole.TELECALLER, OfficeUser.is_active.is_(True), OfficeUser.is_locked.is_(False))
    )
    active_callers = (await db.execute(callers_stmt)).scalars().all()
    if not active_callers:
        raise HTTPException(status_code=400, detail="No active Telecallers available in the system.")

    # Fetch all unassigned NEW leads
    leads_stmt = (
        select(OfficeLead)
        .where(OfficeLead.status == LeadStatus.NEW, OfficeLead.assigned_caller_id.is_(None))
        .order_by(desc(OfficeLead.tier), OfficeLead.created_at.asc())
    )
    unassigned_leads = (await db.execute(leads_stmt)).scalars().all()
    if not unassigned_leads:
        return RoutingResult(total_unassigned=0, total_routed=0, active_callers_count=len(active_callers), allocations={})

    # Calculate current pending load per telecaller
    load_counts = {}
    for caller in active_callers:
        cnt = await db.scalar(
            select(func.count(OfficeLead.id)).where(
                OfficeLead.assigned_caller_id == caller.id,
                OfficeLead.status.in_([LeadStatus.ROUTED_TO_CALLER, LeadStatus.CONTACTED]),
            )
        )
        load_counts[caller.id] = cnt or 0

    allocations = {str(c.emp_code): 0 for c in active_callers}
    routed_total = 0

    for lead in unassigned_leads:
        # Select caller with least load
        target_caller = min(active_callers, key=lambda c: load_counts[c.id])
        lead.assigned_caller_id = target_caller.id
        lead.status = LeadStatus.ROUTED_TO_CALLER
        load_counts[target_caller.id] += 1
        allocations[str(target_caller.emp_code)] += 1
        routed_total += 1

    await db.commit()
    return RoutingResult(
        total_unassigned=len(unassigned_leads),
        total_routed=routed_total,
        active_callers_count=len(active_callers),
        allocations=allocations,
    )

# 4. List Leads with Dynamic Query Matrix & Scoping
@router.get("", response_model=list[LeadResponse])
async def list_leads(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[OfficeUser, Depends(get_current_user)],
    tier: LeadTier | None = Query(None),
    lead_status: LeadStatus | None = Query(None, alias="status"),
    commercial_model: CommercialModel | None = Query(None),
    search: str | None = Query(None),
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
):
    stmt = select(OfficeLead)

    # Scoping Rule: Telecaller can only view their assigned leads
    if current_user.role == OfficeRole.TELECALLER:
        if not current_user.employee_profile:
            return []
        stmt = stmt.where(OfficeLead.assigned_caller_id == current_user.employee_profile.id)

    # Scoping Rule: BDM views assigned leads and subordinates' leads
    elif current_user.role == OfficeRole.BDM:
        if not current_user.employee_profile:
            return []
        stmt = stmt.where(
            or_(
                OfficeLead.assigned_bdm_id == current_user.employee_profile.id,
                OfficeLead.assigned_caller_id.in_(
                    select(OfficeEmployee.id).where(OfficeEmployee.reporting_to_id == current_user.employee_profile.id)
                ),
            )
        )

    if tier:
        stmt = stmt.where(OfficeLead.tier == tier)
    if lead_status:
        stmt = stmt.where(OfficeLead.status == lead_status)
    if commercial_model:
        stmt = stmt.where(OfficeLead.commercial_model == commercial_model)
    if search:
        term = f"%{search.strip().lower()}%"
        stmt = stmt.where(
            or_(
                OfficeLead.company_name.ilike(term),
                OfficeLead.contact_name.ilike(term),
                OfficeLead.email.ilike(term),
                OfficeLead.phone.ilike(term),
                OfficeLead.lead_code.ilike(term),
            )
        )

    stmt = stmt.order_by(desc(OfficeLead.created_at)).offset(offset).limit(limit)
    result = await db.execute(stmt)
    return result.scalars().all()

# 5. Lead Detail
@router.get("/{lead_id}", response_model=LeadResponse)
async def get_lead_dossier(
    lead_id: uuid.UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[OfficeUser, Depends(get_current_user)],
):
    lead = await db.scalar(select(OfficeLead).where(OfficeLead.id == lead_id))
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found.")

    if current_user.role == OfficeRole.TELECALLER:
        if not current_user.employee_profile or lead.assigned_caller_id != current_user.employee_profile.id:
            raise HTTPException(status_code=403, detail="Unauthorized access to this lead record.")

    return lead

# 6. Update Lead Info (Full)
@router.put("/{lead_id}", response_model=LeadResponse)
async def update_lead(
    lead_id: uuid.UUID,
    payload: LeadUpdate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[OfficeUser, Depends(get_current_user)],
):
    lead = await db.scalar(select(OfficeLead).where(OfficeLead.id == lead_id))
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found.")

    for field, value in payload.model_dump(exclude_unset=True).items():
        setattr(lead, field, value)

    # Recalculate dedup hash if phone or email was mutated
    lead.dedup_hash = generate_dedup_hash(lead.email, lead.phone)
    await db.commit()
    await db.refresh(lead)
    return lead

# 7. Granular Status Patch
@router.patch("/{lead_id}/status", response_model=LeadResponse)
async def patch_lead_status(
    lead_id: uuid.UUID,
    payload: LeadPatchStatus,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[OfficeUser, Depends(get_current_user)],
):
    lead = await db.scalar(select(OfficeLead).where(OfficeLead.id == lead_id))
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found.")

    lead.status = payload.status
    if current_user.role == OfficeRole.TELECALLER:
        lead.caller_remarks = payload.remarks
    elif current_user.role == OfficeRole.BDM:
        lead.bdm_remarks = payload.remarks

    await db.commit()
    await db.refresh(lead)
    return lead

# 8. Delete Lead (Archival / Hard purge Admin only)
@router.delete("/{lead_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_lead(
    lead_id: uuid.UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[OfficeUser, Depends(require_roles([OfficeRole.SUPER_ADMIN, OfficeRole.COO]))],
):
    lead = await db.scalar(select(OfficeLead).where(OfficeLead.id == lead_id))
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found.")

    await db.delete(lead)
    await db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)

# 9. Download Streamable Excel Ledger (.xlsx)
@router.get("/export/excel")
async def export_leads_excel(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[OfficeUser, Depends(require_roles([OfficeRole.SUPER_ADMIN, OfficeRole.COO, OfficeRole.OPERATIONS_HEAD]))],
):
    stmt = select(OfficeLead).order_by(desc(OfficeLead.created_at))
    leads = (await db.execute(stmt)).scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Office Leads Ledger"

    headers = [
        "Lead Code", "Company Name", "Contact Name", "Email", "Phone",
        "Country", "State", "Model", "Tier", "Score", "Status",
        "Volume Estimate", "Caller Remarks", "BDM Remarks", "Created At"
    ]
    ws.append(headers)

    for row in ws.iter_cols(min_row=1, max_row=1):
        for cell in row:
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="14120E", end_color="14120E", fill_type="solid")

    for l in leads:
        ws.append([
            l.lead_code, l.company_name, l.contact_name, l.email, l.phone or "",
            l.country, l.state or "", l.commercial_model.value, l.tier.value, l.score,
            l.status.value, l.volume_estimate or "", l.caller_remarks or "",
            l.bdm_remarks or "", l.created_at.strftime("%Y-%m-%d %H:%M")
        ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"Bacchus_Leads_Ledger_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

# 10. Generate Printable PDF Dossier
@router.get("/{lead_id}/download-pdf")
async def download_lead_pdf(
    lead_id: uuid.UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[OfficeUser, Depends(get_current_user)],
):
    lead = await db.scalar(select(OfficeLead).where(OfficeLead.id == lead_id))
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found.")

    pdf_bytes = generate_lead_dossier_pdf(lead)
    filename = f"Lead_Dossier_{lead.lead_code}.pdf"

    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )